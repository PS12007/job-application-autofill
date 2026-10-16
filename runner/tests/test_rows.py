"""Row selection, duplicates / closed postings before opening, and the .xlsx sheet."""
from argparse import Namespace

import openpyxl

import apply
import precheck
import status as S
from config import DEFAULTS
from sheet import XlsxSheet

WD = "https://acme.wd5.myworkdayjobs.com/Careers/job/Toronto/Dev_R{}"


def args(**kw):
    base = dict(start=None, count=None, force=False, redo_failed=False)
    return Namespace(**{**base, **kw})


def make_sheet(tmp_path, rows):
    f = tmp_path / "apps.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Row ID", "Company", "Role", "Link", "Status"])  # Notes / Last Updated missing on purpose
    for r in rows:
        ws.append(r)
    ws["D3"].hyperlink = WD.format(2)  # hyperlink target wins over the cell text
    ws["D3"].value = "link"
    wb.save(f)
    return f, XlsxSheet(f, "", dry_run=False)


ROWS = [
    ["1", "Acme", "Dev", WD.format(1), ""],
    ["2", "Acme", "Dev 2", "x", "queued"],
    ["3", "Acme", "Dev", WD.format(1).replace("/Careers/", "/en-US/Careers/"), ""],  # same job as row 2
    ["4", "Beta", "Dev", WD.format(4), "Submitted"],
    ["5", "Beta", "Dev", WD.format(4) + "/apply", ""],  # same job as the submitted row 5
    ["6", "Gamma", "Dev", WD.format(6), "Failed"],
    ["7", "Gamma", "Dev", WD.format(7), "Needs you: unanswered questions"],
    ["8", "Gamma", "Dev", '=HYPERLINK("' + WD.format(8) + '", "Apply")', "my own note"],
    ["9", "Gamma", "Dev", "not a url", ""],
]


def test_sheet_reads_links_and_adds_columns(tmp_path):
    f, sh = make_sheet(tmp_path, ROWS)
    assert [r.link for r in sh.rows][:3] == [WD.format(1), WD.format(2), ROWS[2][3]]
    assert sh.rows[7].link == WD.format(8)
    assert sh.rows[1].status == S.QUEUED  # canonical case
    sh.update(sh.rows[0], S.READY, "note")
    sh.close()
    ws = openpyxl.load_workbook(f).active
    assert [c.value for c in ws[1]] == ["Row ID", "Company", "Role", "Link", "Status", "Notes", "Last Updated"]
    assert ws["E2"].value == S.READY and ws["F2"].value == "note" and ws["G2"].value


def test_select_rows(tmp_path):
    _, sh = make_sheet(tmp_path, ROWS)
    nums = lambda rows: [r.num for r in rows]
    assert nums(apply.select_rows(sh.rows, args())) == [2, 3, 4, 6, 8]
    assert nums(apply.select_rows(sh.rows, args(redo_failed=True))) == [2, 3, 4, 6, 7, 8]
    assert nums(apply.select_rows(sh.rows, args(force=True))) == [2, 3, 4, 5, 6, 7, 8, 9]
    assert nums(apply.select_rows(sh.rows, args(start=5))) == [6, 8]


def test_prepare_marks_duplicates_and_closed(tmp_path, monkeypatch):
    _, sh = make_sheet(tmp_path, ROWS)
    infos = {WD.format(7): precheck.JobInfo(closed=True, req_id="R7")}
    monkeypatch.setattr(precheck, "lookup_many", lambda urls: {u: infos.get(u) for u in urls})
    cfg = {"run": {**DEFAULTS["run"], "precheck": True}}
    queue, skipped = apply.prepare(apply.select_rows(sh.rows, args()), sh, cfg, args())
    assert [r.num for r in queue] == [2, 3]
    by = {r.num: r for r in sh.rows}
    assert by[4].status == S.DUPLICATE and "row 2" in by[4].notes
    assert by[6].status == S.DUPLICATE and "row 5 (Submitted)" in by[6].notes
    assert by[8].status == S.CLOSED
    assert [r.num for r in skipped] == [4, 6, 8]


def test_prepare_count_leaves_later_rows_alone(tmp_path, monkeypatch):
    _, sh = make_sheet(tmp_path, ROWS)
    monkeypatch.setattr(precheck, "lookup_many", lambda urls: {u: precheck.JobInfo(closed=True) for u in urls})
    cfg = {"run": {**DEFAULTS["run"], "precheck": False}}
    queue, _ = apply.prepare(apply.select_rows(sh.rows, args()), sh, cfg, args(count=1))
    assert [r.num for r in queue] == [2]
    assert {r.num: r.status for r in sh.rows}[8] == S.NEEDS_ANSWERS


def test_canonical():
    assert S.canonical("  ready TO submit ") == S.READY
    assert S.canonical("") == ""
    assert S.canonical("Waiting on recruiter") == "Waiting on recruiter"
