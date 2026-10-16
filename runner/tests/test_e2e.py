"""End-to-end: the real runner + extension in a visible Chromium against the mock Workday
site (tests/mock). Slow (a few minutes); run with:  pytest --e2e

What it checks:
- sign in, "Sign in with email", create account (only the required box ticked, the
  honeypot left empty), "Use My Last Application", verify email via the (fake) inbox
- statuses for ready / unanswered questions / closed / already applied / CAPTCHA / duplicate
- the watcher: Submitted after you submit, Ready after you answer and reach Review
- questions.yaml collects the unanswered questions; answering there fills them next run
"""
import json
import os
import socket
import subprocess
import sys
import threading
import time
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import openpyxl
import pytest
import yaml

import status as S

pytestmark = pytest.mark.e2e
HERE = Path(__file__).resolve().parent
RUNNER = HERE.parent


def free_port():
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


class Quiet(SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass


@pytest.fixture(scope="module")
def mock_site():
    port = free_port()
    srv = ThreadingHTTPServer(("127.0.0.1", port), partial(Quiet, directory=str(HERE / "mock")))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield port
    srv.shutdown()


U = lambda tenant, job="ready": f"http://{tenant}.wd5.myworkdayjobs.com/?job={job}"
ROWS = {
    2: ("Acme", U("acme")),
    3: ("Beta", U("beta", "questions")),
    4: ("Gamma", U("gamma")),
    5: ("Acme closed", U("acme", "closed")),
    6: ("Acme applied", U("acme", "applied")),
    7: ("Acme captcha", U("acme", "captcha")),
    8: ("Acme again", U("acme")),
    9: ("Delta", U("delta")),
    10: ("Delta earlier", U("delta", "older")),  # already Submitted: "Use My Last Application" at delta
    11: ("Epsilon", U("epsilon")),  # back on Sign In after creating the account: verify by email
    12: ("Epsilon 2", U("epsilon", "second")),  # same company: no second account
}


def setup(tmp, port):
    debug = free_port()
    sheet = tmp / "apps.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Row ID", "Company", "Role", "Link", "Status", "Notes", "Last Updated"])
    for n, (company, link) in ROWS.items():
        wb.active.append([str(n), company, "Software Developer", link, S.SUBMITTED if n == 10 else ""])
    wb.save(sheet)
    cfg = {
        "sheet": {"backend": "xlsx", "xlsx_path": str(sheet)},
        "browser": {"user_data_dir": str(tmp / "browser"), "debug_port": debug,
                    "extra_args": [f"--host-resolver-rules=MAP *.myworkdayjobs.com 127.0.0.1:{port}, "
                                   f"MAP *.hcaptcha.com 127.0.0.1:{port}"]},
        "run": {"max_tabs": 10, "app_timeout_s": 120, "delay_between_s": 1, "precheck": False},
        "notify": {"sound": False, "desktop": False},
    }
    (tmp / "config.yaml").write_text(yaml.safe_dump(cfg), encoding="utf-8")
    profile = json.loads((RUNNER / "profile.example.json").read_text(encoding="utf-8"))
    (tmp / "resume.pdf").write_bytes(b"%PDF-1.4 mock resume")
    profile["resumePath"] = str(tmp / "resume.pdf")
    (tmp / "profile.json").write_text(json.dumps(profile), encoding="utf-8")
    (tmp / "resumes.yaml").write_text("tracks: {}\n", encoding="utf-8")
    return sheet, debug


def start(tmp, *extra):
    env = {k: v for k, v in os.environ.items() if not k.startswith("IMAP_")}
    env["WORKDAY_PASSWORD"] = "pw-test"
    args = [sys.executable, str(HERE / "e2e_run.py"), "--config", str(tmp / "config.yaml"),
            "--profile", str(tmp / "profile.json"), "--resumes", str(tmp / "resumes.yaml"),
            "--questions", str(tmp / "questions.yaml"), *extra]
    log = open(tmp / f"run_{int(time.time())}.out", "w", encoding="utf-8")
    return subprocess.Popen(args, cwd=RUNNER, env=env, stdout=log, stderr=subprocess.STDOUT), log


def statuses(sheet):
    for _ in range(10):
        try:
            ws = openpyxl.load_workbook(sheet).active
            return {r: (ws.cell(r, 5).value or "", ws.cell(r, 6).value or "") for r in range(2, ws.max_row + 1)}
        except Exception:
            time.sleep(0.5)
    raise RuntimeError("couldn't read the sheet")


def wait_for(sheet, want, proc, timeout=300):
    end = time.time() + timeout
    got = {}
    while time.time() < end:
        got = statuses(sheet)
        if all(got.get(r, ("",))[0] == s for r, s in want.items()):
            return got
        if proc.poll() is not None:
            break
        time.sleep(2)
    pytest.fail(f"expected {want}, sheet has { {r: got.get(r) for r in want} }")


def close_job_tabs(ctx):
    for p in list(ctx.pages):
        if "myworkdayjobs" in p.url:
            p.close()


def test_full_run(mock_site, tmp_path):
    from playwright.sync_api import sync_playwright

    sheet, debug = setup(tmp_path, mock_site)
    proc, log = start(tmp_path)
    try:
        got = wait_for(sheet, {
            2: S.READY, 3: S.NEEDS_ANSWERS, 4: S.READY, 5: S.CLOSED, 6: S.ALREADY,
            7: S.NEEDS_CAPTCHA, 8: S.DUPLICATE, 9: S.READY, 11: S.NEEDS_VERIFY, 12: S.NEEDS_VERIFY,
        }, proc)
        assert "row 2" in got[8][1]
        assert "earlier row" in got[12][1]
        assert "favourite mock colour" in got[3][1]

        bank = yaml.safe_load((tmp_path / "questions.yaml").read_text(encoding="utf-8"))
        by_q = {i["question"].rstrip("*"): i for i in bank}
        assert by_q["Are you willing to work mock weekends?"]["options"] == ["Yes", "No"]
        assert by_q["What is your favourite mock colour?"]["required"]

        with sync_playwright() as pw:
            ctx = pw.chromium.connect_over_cdp(f"http://127.0.0.1:{debug}").contexts[0]
            tab = lambda s: next(p for p in ctx.pages if s in p.url)
            beta = tab("beta.")
            assert beta.evaluate("[localStorage.terms, localStorage.marketing, localStorage.bee || '']") == ["true", "false", ""]
            assert tab("delta.").evaluate("localStorage.usedLast") == "1"
            assert tab("epsilon.").evaluate("localStorage.created") == "1"  # one account per company

            # You verify epsilon from your email and sign in in its first tab: the watcher
            # fills it, then retries the other epsilon tab, which now signs in by itself.
            eps = tab("epsilon.wd5.myworkdayjobs.com/?job=ready")
            eps.evaluate("localStorage.setItem('verified', '1'); sessionStorage.setItem('in', '1'); S.view = 'form'; render()")
            wait_for(sheet, {11: S.READY, 12: S.READY}, proc, timeout=90)
            assert beta.evaluate("localStorage.usedLast || ''") == ""  # never applied at beta: Apply Manually

            # You submit the acme tab: the watcher marks it Submitted and closes it.
            tab("acme.wd5.myworkdayjobs.com/?job=ready").click('[data-automation-id="bottom-navigation-next-button"]')
            # You answer beta's questions and go to Review: Ready to submit.
            beta.fill("#q1", "Blue")
            beta.check('input[name="q2"][value="y"]')
            beta.click('[data-automation-id="bottom-navigation-next-button"]')
            wait_for(sheet, {2: S.SUBMITTED, 3: S.READY}, proc, timeout=60)

            close_job_tabs(ctx)
        assert proc.wait(timeout=60) == 0
    finally:
        if proc.poll() is None:
            proc.kill()
        log.close()

    # Next run: answers typed into questions.yaml fill the same questions automatically.
    bank = yaml.safe_load((tmp_path / "questions.yaml").read_text(encoding="utf-8"))
    for i in bank:
        i["answer"] = "Blue" if "colour" in i["question"] else "Yes"
    (tmp_path / "questions.yaml").write_text(yaml.safe_dump(bank), encoding="utf-8")
    wb = openpyxl.load_workbook(sheet)
    wb.active.cell(3, 5).value = S.QUEUED
    wb.save(sheet)

    proc, log = start(tmp_path, "--start", "3", "--count", "1")
    try:
        got = wait_for(sheet, {3: S.READY}, proc, timeout=180)
        with sync_playwright() as pw:
            ctx = pw.chromium.connect_over_cdp(f"http://127.0.0.1:{debug}").contexts[0]
            beta = next(p for p in ctx.pages if "beta." in p.url)
            assert beta.inner_text("h2") == "Review"  # both questions answered from questions.yaml
            close_job_tabs(ctx)
        assert proc.wait(timeout=60) == 0
    finally:
        if proc.poll() is None:
            proc.kill()
        log.close()
    assert "Using 2 answer(s) from questions.yaml" in Path(log.name).read_text(encoding="utf-8")
