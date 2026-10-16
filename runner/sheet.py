"""sheet.py — the application queue: a local .xlsx (openpyxl) or a Google Sheet (gspread).

Reads Link (plus Row ID / Company / Role for messages) from row 2 down, with headers in
row 1. Writes only Status, Notes and Last Updated, and adds those headers if missing.
"""
import logging
import re
import time
from dataclasses import dataclass
from datetime import datetime

import status as S

log = logging.getLogger("runner")

READ_COLS = {"row_id": "Row ID", "company": "Company", "role": "Role", "link": "Link", "resume": "Resume"}
WRITE_COLS = {"status": "Status", "notes": "Notes", "updated": "Last Updated"}
HEADERS = ["Row ID", "Company", "Role", "Link", "Status", "Notes", "Last Updated"]


@dataclass
class Row:
    num: int  # spreadsheet row number (2 = first data row)
    row_id: str
    company: str
    role: str
    link: str
    status: str
    notes: str
    resume: str = ""  # optional "Resume" column: a track name or file path

    @property
    def name(self):
        return " / ".join(x for x in (self.company, self.role) if x) or self.link[:60]


def _norm(s):
    return " ".join(str(s or "").split()).lower()


_HYPERLINK = re.compile(r'^=\s*HYPERLINK\(\s*"([^"]+)"', re.I)


def _link_text(value):
    """Plain URL from a cell value, including =HYPERLINK("url", "text") formulas."""
    s = str(value or "").strip()
    m = _HYPERLINK.match(s)
    return m.group(1) if m else s


class _Base:
    def __init__(self, dry_run):
        self.dry_run = dry_run
        self.cols = {}
        self.rows = []

    def _map_header(self, header):
        """Find columns by header name; returns the write-columns that are missing."""
        for i, h in enumerate(header):
            for key, name in {**READ_COLS, **WRITE_COLS}.items():
                if _norm(h) == _norm(name) and key not in self.cols:
                    self.cols[key] = i + 1
        if "link" not in self.cols:
            raise SystemExit("The sheet has no 'Link' column in row 1.")
        last = max((i + 1 for i, h in enumerate(header) if _norm(h)), default=0)
        missing = [k for k in WRITE_COLS if k not in self.cols]
        for n, k in enumerate(missing):
            self.cols[k] = last + 1 + n
        return missing

    def _make_row(self, num, get):
        val = lambda k: str(get(self.cols[k]) or "").strip() if k in self.cols else ""
        return Row(num, val("row_id"), val("company"), val("role"), val("link"),
                   S.canonical(val("status")), val("notes"), val("resume"))

    def update(self, row, status, notes=""):
        """Set Status / Notes / Last Updated for one row (written immediately)."""
        row.status, row.notes = status, notes
        if self.dry_run:
            log.info("[dry run] row %s would be set to %s%s", row.num, status, f" ({notes})" if notes else "")
            return
        stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
        self._write(row.num, {"status": status, "notes": notes, "updated": stamp})

    def close(self):
        pass


class XlsxSheet(_Base):
    def __init__(self, file, worksheet, dry_run):
        import openpyxl

        super().__init__(dry_run)
        self.path = file
        self._dirty = False
        self._warned = False
        if not file.exists():
            if dry_run:
                raise SystemExit(f"{file} doesn't exist yet. Run once without --dry-run to create it.")
            wb = openpyxl.Workbook()
            wb.active.title = worksheet or "Applications"
            wb.active.append(HEADERS)
            wb.save(file)
            log.info("Created %s with the expected headers. Add your links and run again.", file)
        self.wb = openpyxl.load_workbook(file)
        self.ws = self.wb[worksheet] if worksheet else self.wb.worksheets[0]

        header = [c.value for c in self.ws[1]]
        missing = self._map_header(header)
        if missing and not dry_run:
            for k in missing:
                self.ws.cell(1, self.cols[k], WRITE_COLS[k])
            self._save()
            log.info("Added column(s) to the sheet: %s", ", ".join(WRITE_COLS[k] for k in missing))

        link_col = self.cols["link"]
        for r in range(2, self.ws.max_row + 1):
            row = self._make_row(r, lambda c: self.ws.cell(r, c).value)
            cell = self.ws.cell(r, link_col)
            row.link = (cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else _link_text(cell.value))
            self.rows.append(row)

    def _write(self, num, vals):
        for k, v in vals.items():
            self.ws.cell(num, self.cols[k], v)
        self._dirty = True
        self._save()

    def _save(self):
        try:
            self.wb.save(self.path)
            self._dirty = False
            if self._warned:
                log.info("Saved %s again.", self.path.name)
            self._warned = False
        except PermissionError:
            if not self._warned:
                log.warning("Can't save %s (is it open in Excel?). Close it; changes are kept "
                            "and saved on the next update.", self.path.name)
            self._warned = True

    def close(self):
        for _ in range(15):
            if not self._dirty:
                return
            self._save()
            if self._dirty:
                time.sleep(2)
        if self._dirty:
            log.error("Could not save %s. Statuses from this run were NOT written.", self.path)


class GoogleSheet(_Base):
    def __init__(self, sheet_id, worksheet, sa_file, dry_run):
        import gspread

        super().__init__(dry_run)
        if not sheet_id:
            raise SystemExit("Set sheet.gsheets_id in config.yaml.")
        if not sa_file.exists():
            raise SystemExit(f"Service account key not found: {sa_file} (see README).")
        gc = gspread.service_account(filename=str(sa_file))
        try:
            sh = gc.open_by_url(sheet_id) if sheet_id.startswith("http") else gc.open_by_key(sheet_id)
        except gspread.SpreadsheetNotFound:
            raise SystemExit("Sheet not found. Share it (Editor) with the service account's email (see README).")
        self.ws = sh.worksheet(worksheet) if worksheet else sh.sheet1

        values = self.ws.get_all_values()
        header = values[0] if values else []
        missing = self._map_header(header)
        if missing and not dry_run:
            need = max(self.cols[k] for k in missing)
            if need > self.ws.col_count:
                self.ws.add_cols(need - self.ws.col_count)
            self._batch([(1, self.cols[k], WRITE_COLS[k]) for k in missing])
            log.info("Added column(s) to the sheet: %s", ", ".join(WRITE_COLS[k] for k in missing))

        for i, vals in enumerate(values[1:], start=2):
            row = self._make_row(i, lambda c: vals[c - 1] if c - 1 < len(vals) else "")
            row.link = _link_text(row.link)
            self.rows.append(row)

    def _batch(self, cells):
        from gspread.exceptions import APIError
        from gspread.utils import rowcol_to_a1

        data = [{"range": rowcol_to_a1(r, c), "values": [[v]]} for r, c, v in cells]
        for attempt in range(5):
            try:
                self.ws.batch_update(data, raw=True)
                return
            except APIError as e:  # 429 rate limit / transient 5xx
                code = getattr(getattr(e, "response", None), "status_code", 0)
                if code not in (429, 500, 502, 503) or attempt == 4:
                    log.error("Google Sheets write failed: %s", e)
                    return
                time.sleep(5 * (attempt + 1))

    def _write(self, num, vals):
        self._batch([(num, self.cols[k], v) for k, v in vals.items()])


def open_sheet(cfg, dry_run):
    from config import path

    sc = cfg["sheet"]
    if sc["backend"] == "gsheets":
        return GoogleSheet(str(sc["gsheets_id"]).strip(), sc["worksheet"], path(sc["service_account_file"]), dry_run)
    if sc["backend"] == "xlsx":
        return XlsxSheet(path(sc["xlsx_path"]), sc["worksheet"], dry_run)
    raise SystemExit(f"sheet.backend must be xlsx or gsheets, not {sc['backend']!r}")
